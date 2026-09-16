from __future__ import annotations

import json
import os
from collections import defaultdict
from datetime import date
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
BACKEND = ROOT / "backend"
OUTPUT_DIR = ROOT / "docs" / "deliverables" / "system_docs_20260729"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

os.environ.setdefault("PYTHONPATH", str(BACKEND))

from app.main import app  # noqa: E402
from app.db.session import Base  # noqa: E402
from app import models as _models  # noqa: E402,F401


DOC_DATE = date(2026, 7, 29)
VERSION = "V1.0"
SYSTEM_NAME = "Star Rail Agents（星穹智库）"

NAVY = RGBColor(31, 77, 120)
BLUE = RGBColor(46, 116, 181)
GRAY = RGBColor(89, 89, 89)
LIGHT = "F2F4F7"
PALE_BLUE = "EAF2F8"
WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(0, 0, 0)
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS = {"top": 80, "bottom": 80, "start": 120, "end": 120}


def set_run_font(
    run,
    *,
    size: float | None = None,
    bold: bool | None = None,
    color: RGBColor | None = None,
    east_asia: str = "Microsoft YaHei",
    latin: str = "Calibri",
) -> None:
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in CELL_MARGINS.items():
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_begin, instr_text, fld_char_end])
    set_run_font(run, size=9, color=GRAY)


def configure_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.42)
    section.footer_distance = Inches(0.42)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    heading_tokens = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, NAVY, 8, 4),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167


def setup_header_footer(doc: Document, short_title: str) -> None:
    for section in doc.sections:
        header = section.header
        p = header.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(f"{SYSTEM_NAME}  |  {short_title}")
        set_run_font(run, size=8.5, color=GRAY)
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.paragraph_format.space_after = Pt(0)
        run = fp.add_run(f"{VERSION}  ·  {DOC_DATE.isoformat()}  ·  第 ")
        set_run_font(run, size=9, color=GRAY)
        add_page_field(fp)
        run = fp.add_run(" 页")
        set_run_font(run, size=9, color=GRAY)


def add_cover(
    doc: Document,
    title: str,
    subtitle: str,
    doc_code: str,
    audience: str = "项目评审、开发、测试与运维人员",
) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(90)
    p.paragraph_format.space_after = Pt(16)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("系统工程文档")
    set_run_font(run, size=11, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    set_run_font(run, size=28, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(54)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(subtitle)
    set_run_font(run, size=14, color=GRAY)

    rows = [
        ("项目名称", SYSTEM_NAME),
        ("文档编号", doc_code),
        ("版本 / 日期", f"{VERSION} / {DOC_DATE.isoformat()}"),
        ("适用对象", audience),
        ("编制依据", "当前仓库源码、配置、数据库迁移与 FastAPI OpenAPI 定义"),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    set_table_geometry(table, [2300, 7060])
    set_repeat_table_header(table.rows[0])
    for i, (label, value) in enumerate(rows):
        table.cell(i, 0).text = label
        table.cell(i, 1).text = value
        set_cell_shading(table.cell(i, 0), PALE_BLUE)
        for run in table.cell(i, 0).paragraphs[0].runs:
            set_run_font(run, size=10, bold=True, color=NAVY)
        for run in table.cell(i, 1).paragraphs[0].runs:
            set_run_font(run, size=10, color=BLACK)
    doc.add_page_break()


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_para(doc: Document, text: str, *, bold_lead: str | None = None) -> None:
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, bold=True)
        rest = p.add_run(text[len(bold_lead) :])
        set_run_font(rest)
    else:
        run = p.add_run(text)
        set_run_font(run)


def add_bullets(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run_font(run)


def add_numbered(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        run = p.add_run(item)
        set_run_font(run)


def add_callout(doc: Document, label: str, text: str, fill: str = PALE_BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [TABLE_WIDTH_DXA])
    set_repeat_table_header(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(f"{label}：")
    set_run_font(run, bold=True, color=NAVY)
    run = p.add_run(text)
    set_run_font(run)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[Any]],
    widths_dxa: list[int],
    *,
    font_size: float = 9.3,
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths_dxa)
    set_repeat_table_header(table.rows[0])
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = header
        set_cell_shading(cell, LIGHT)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cell.paragraphs[0].runs:
            set_run_font(run, size=font_size, bold=True, color=NAVY)
    for values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cells[idx].text = str(value)
            for p in cells[idx].paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    set_run_font(run, size=font_size)
    set_table_geometry(table, widths_dxa)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def new_doc(title: str, subtitle: str, code: str, short_title: str) -> Document:
    doc = Document()
    configure_styles(doc)
    setup_header_footer(doc, short_title)
    add_cover(doc, title, subtitle, code)
    return doc


def save_doc(doc: Document, filename: str) -> Path:
    path = OUTPUT_DIR / filename
    doc.core_properties.title = path.stem
    doc.core_properties.subject = SYSTEM_NAME
    doc.core_properties.author = "Star Rail Agents 项目组"
    doc.core_properties.keywords = "可行性,需求,设计,接口,数据库,FastAPI,Vue,PostgreSQL,Milvus"
    doc.save(path)
    return path


def build_feasibility() -> Path:
    doc = new_doc(
        "可行性分析报告",
        "经济、操作、技术与运行可行性评估",
        "SRA-FEA-001",
        "可行性分析",
    )
    add_callout(
        doc,
        "评估结论",
        "基于当前已实现原型与开源技术栈，项目具备毕业设计/教学演示和小规模试运行可行性。"
        "若转为公网生产服务，应补充真实并发压测、云资源报价、内容合规流程和灾备演练后再作投产决策。",
    )

    add_heading(doc, "1. 项目概述")
    add_para(
        doc,
        "本项目面向《崩坏：星穹铁道》玩家，提供官方资料检索、带引用的智能问答、角色养成计算、"
        "配队推荐、每周规划、会话与长期记忆、自定义角色创作及社区审核等能力。系统采用 Vue 3 "
        "前端、FastAPI 后端、PostgreSQL 业务库、Milvus 向量库、BGE 嵌入与重排服务以及可配置的大模型提供商。"
    )
    add_heading(doc, "2. 评估边界与假设")
    add_bullets(
        doc,
        [
            "结论基于 2026-07-29 仓库中的源码、Docker Compose、Alembic 迁移和测试结构。",
            "经济数据采用成本类别和估算口径，不代表已经发生的采购或云服务账单。",
            "性能与可用性数字属于建议验收目标；未见正式压测报告的项目，不应表述为已达标。",
            "游戏内容及商标权利归相关权利人所有；面向公众运营时需另行完成授权、版权与平台条款审查。",
        ],
    )

    add_heading(doc, "3. 经济可行性")
    add_para(
        doc,
        "核心框架与基础设施均可使用开源版本，软件许可成本低。主要可变成本来自计算资源、"
        "大模型 API 调用、GPU 推理、存储备份、域名与运维时间。教学演示可复用开发机并使用 mock "
        "模型或本地模型，显著降低现金支出；生产环境则需按日活、对话量、向量规模和响应时延估算。",
    )
    add_table(
        doc,
        ["成本项", "原型/教学阶段", "生产阶段主要变量", "控制措施"],
        [
            ["软件许可", "Vue、FastAPI、PostgreSQL、Milvus 等可零许可费使用", "需核对各依赖及模型许可证", "建立依赖清单与许可证审计"],
            ["计算资源", "可使用现有开发机；GPU 可选", "CPU/GPU 规格、并发、运行时长", "模型服务独立伸缩，支持 CPU 降级"],
            ["大模型调用", "mock 模式可不产生外部调用费", "请求量、上下文长度、模型单价", "路由分级、缓存、限额与成本监控"],
            ["存储与备份", "本地命名卷即可", "用户数据、向量索引、图片与保留周期", "冷热分层、定期备份和生命周期策略"],
            ["人员投入", "以课程/毕业设计工时为主", "研发、内容维护、审核、值守", "自动化测试、迁移与审核工作台"],
        ],
        [1650, 2500, 2500, 2710],
    )
    add_para(
        doc,
        "建议的预算测算公式：月度总成本 = 基础主机与数据库 + GPU/模型服务 + LLM 调用量 × 单价 + "
        "对象存储与流量 + 备份 + 运维人力。上线前至少以低、中、高三档业务量报价，并设置单用户和全局调用上限。"
    )
    add_callout(doc, "经济性判断", "原型和小规模试运行经济可行；公开生产运营为“有条件可行”，需结合真实流量完成报价与预算审批。")

    add_heading(doc, "4. 操作可行性")
    add_para(
        doc,
        "玩家通过网页完成登录、资料浏览、智能问答、角色池登记、养成与队伍管理；管理员通过审核页面处理"
        "自定义角色和活动攻略。交互路径与既有游戏工具使用习惯相近，学习成本可控。敏感写操作由用户显式触发，"
        "智能体不会自动写入长期记忆或公开玩家内容，降低误操作风险。"
    )
    add_table(
        doc,
        ["角色", "典型任务", "操作条件", "可行性要点"],
        [
            ["游客", "浏览公开目录、活动、剧情与社区内容", "浏览器", "无需安装客户端；入口清晰"],
            ["注册用户", "对话、角色池、记忆、队伍、养成计划", "登录并持有 JWT", "状态可保存；资源按用户隔离"],
            ["创作者", "四阶段角色创作、提交审核", "注册用户", "草稿、版本、审核状态分离"],
            ["管理员", "审核、下架、恢复公开内容", "管理员账号", "权限校验和操作记录降低治理风险"],
        ],
        [1300, 2800, 1900, 3360],
    )
    add_bullets(
        doc,
        [
            "建议补充首次使用引导、空状态说明和危险操作二次确认。",
            "建议按 WCAG 2.1 AA 检查键盘可达性、对比度、焦点顺序和错误提示。",
            "管理员需配置审核准则、申诉处理和违规内容留痕制度。",
        ],
    )
    add_callout(doc, "操作性判断", "现有角色与流程明确，操作可行；无障碍和新手引导应作为上线前验收项。")

    add_heading(doc, "5. 技术可行性")
    add_table(
        doc,
        ["技术域", "现有实现", "可行性依据", "主要风险"],
        [
            ["前端", "Vue 3、TypeScript、Vite、Pinia、Vue Router", "生态成熟，组件化与类型检查完善", "复杂页面状态和移动端适配"],
            ["后端", "FastAPI、Pydantic、SQLAlchemy、Alembic", "自动 OpenAPI、校验与迁移能力完整", "同步数据库访问在高并发下需评估"],
            ["检索", "Milvus、BGE-M3、Reranker", "适合语义检索与重排", "模型显存、索引质量和冷启动"],
            ["智能体", "Router + 专业 Agent + 统一证据协议", "语义任务与确定性计算分离", "外部模型波动、幻觉与成本"],
            ["部署", "Docker Compose、健康检查、命名卷", "环境可复制，服务依赖明确", "单机编排不等同高可用"],
            ["测试", "pytest、ruff、前端 type-check/build", "已有单元与集成测试基础", "缺少正式压测和端到端浏览器回归证据"],
        ],
        [1300, 2400, 2500, 3160],
    )
    add_para(
        doc,
        "技术上不存在必须依赖自研基础设施的阻断项。官方资料与用户创作采用独立数据路径，RAG 回答经过"
        "Claim-Citation-Validation-Filtering 协议过滤；配队、养成和每周规划优先使用确定性计算，能够降低数值幻觉。"
    )
    add_callout(doc, "技术性判断", "技术路线成熟且实现基础充分，技术可行；生产化重点是压测、限流、缓存、可观测性和模型服务降级。")

    add_heading(doc, "6. 运行可行性")
    add_para(
        doc,
        "Docker Compose 已编排 PostgreSQL、etcd、MinIO、Milvus、BGE 模型服务、后端与前端，"
        "并为核心服务配置健康检查和自动重启。后端启动前执行 Alembic 迁移，空库能够迁移到当前版本。"
    )
    add_table(
        doc,
        ["运行要素", "当前机制", "上线前补强"],
        [
            ["启动与恢复", "depends_on 健康条件、restart unless-stopped", "增加启动超时告警和故障演练"],
            ["数据持久化", "PostgreSQL/Milvus/MinIO/etcd 命名卷", "异机备份、恢复演练、RPO/RTO"],
            ["配置", ".env 与容器环境变量", "密钥管理服务、配置分环境、禁止默认密码"],
            ["监控", "HTTP 健康检查", "日志聚合、指标、追踪、模型耗时和成本看板"],
            ["伸缩", "服务职责已拆分", "生产使用编排平台或托管服务实现副本与滚动升级"],
        ],
        [1600, 3400, 4360],
    )
    add_callout(doc, "运行性判断", "本地与单机试运行可行；生产高可用运行属于“有条件可行”，须建立备份、监控、容量和应急体系。")

    add_heading(doc, "7. 风险与对策")
    add_table(
        doc,
        ["风险", "概率", "影响", "对策", "责任域"],
        [
            ["模型 API 不可用或响应变慢", "中", "高", "超时、重试、熔断、mock/规则降级", "后端/运维"],
            ["BGE 模型或 GPU 资源不足", "中", "中高", "CPU 模式、延迟加载、批量参数与容量监控", "模型服务"],
            ["知识过期或引用不一致", "中", "高", "来源登记、数据审计、版本化重建索引", "数据治理"],
            ["用户内容违规或侵权", "中", "高", "审核、下架、操作留痕、申诉与规则公示", "运营/管理员"],
            ["默认密钥或密码进入生产", "低中", "高", "启动校验、密钥轮换、CI 秘密扫描", "安全/运维"],
            ["单机故障导致中断", "中", "高", "备份、恢复演练、托管数据库与多副本", "运维"],
        ],
        [2400, 850, 850, 3600, 1660],
    )

    add_heading(doc, "8. 综合结论与准入条件")
    add_para(
        doc,
        "项目在经济、操作和技术层面总体可行，在本地及小规模环境具备运行可行性。建议按以下门槛决定是否公开上线："
    )
    add_numbered(
        doc,
        [
            "完成目标并发下的接口压测、模型服务容量测试和关键页面端到端回归。",
            "替换所有默认密码与开发 JWT 密钥，落实限流、审计日志和敏感配置管理。",
            "建立 PostgreSQL 与 Milvus 的备份、恢复验证及明确的 RPO/RTO。",
            "完成内容来源、模型许可证、游戏素材使用和社区审核规则审查。",
            "按预估流量核算云资源及 LLM 月度成本，设定预算告警和降级开关。",
        ],
    )
    return save_doc(doc, "01_可行性分析报告.docx")


def build_requirements() -> Path:
    doc = new_doc(
        "需求分析报告",
        "功能需求、非功能需求与验收基线",
        "SRA-REQ-001",
        "需求分析",
    )
    add_callout(
        doc,
        "文档用途",
        "定义系统当前实现所覆盖的业务边界，并将关键能力转化为可追踪、可测试的需求项。"
        "标注为“目标”的非功能指标需通过专项测试后方可判定达成。",
    )
    add_heading(doc, "1. 背景与目标")
    add_para(
        doc,
        "玩家资料分散、数值养成计算繁琐、攻略质量不一，通用大模型又容易给出无来源结论。"
        "系统目标是将结构化资料、语义检索、确定性计算与多智能体路由整合为统一网页应用，"
        "让玩家获得可引用、可校验、可持续保存的辅助结果。"
    )
    add_heading(doc, "2. 范围")
    add_heading(doc, "2.1 范围内", level=2)
    add_bullets(
        doc,
        [
            "账号注册、登录、当前用户识别与管理员授权。",
            "角色、光锥、遗器、物品、怪物、活动和剧情资料浏览。",
            "智能问答、会话历史、异步聊天任务、长期记忆确认与维护。",
            "角色池、练度、配队、养成方案和每周体力规划。",
            "自定义角色创作、版本化、社区发布、审核与内容治理。",
            "Docker 化部署、数据库迁移、健康检查和 OpenAPI 文档。",
        ],
    )
    add_heading(doc, "2.2 范围外", level=2)
    add_bullets(
        doc,
        [
            "与游戏客户端账号、抽卡记录或实时战斗数据的官方直连。",
            "自动代替玩家执行游戏操作或修改游戏数据。",
            "面向大规模公网生产环境的多地域高可用保证。",
            "对外部模型、游戏版本更新和第三方内容持续可用性的绝对保证。",
        ],
    )

    add_heading(doc, "3. 用户角色与权限")
    add_table(
        doc,
        ["角色", "身份条件", "核心权限", "限制"],
        [
            ["游客", "未登录", "浏览公开资料、公开活动/剧情/社区内容，匿名问答", "不能保存用户数据或提交内容"],
            ["注册用户", "有效 JWT", "管理本人角色池、记忆、队伍、计划、会话和创作草稿", "只能访问本人私有资源"],
            ["创作者", "注册用户", "提交自定义角色或活动攻略供审核", "提交不等于立即公开"],
            ["管理员", "用户名位于管理员配置", "审核、驳回、下架、恢复社区内容", "操作需鉴权并建议留痕"],
            ["外部模型服务", "后端受控调用", "完成语义路由、生成或嵌入/重排", "不得直接写入用户业务数据"],
        ],
        [1250, 1750, 3600, 2760],
    )

    add_heading(doc, "4. 功能需求")
    functional = [
        ("FR-AUTH-01", "账号注册", "游客提交用户名、邮箱、显示名和密码，系统校验唯一性并创建账号。", "合法输入返回 201 和访问令牌；冲突返回 409；非法输入返回 422。"),
        ("FR-AUTH-02", "账号登录", "用户可用账号标识与密码登录。", "认证成功返回 Bearer JWT 和用户信息；失败不泄露密码细节。"),
        ("FR-AUTH-03", "身份与授权", "后端从 Bearer 令牌识别当前用户，并限制私有资源和管理员接口。", "未认证返回 401；越权返回 403/404。"),
        ("FR-CAT-01", "官方目录浏览", "用户可查看角色列表与详情，以及物品、光锥、遗器和怪物等知识实体。", "列表和详情均可访问；不存在资源返回 404。"),
        ("FR-STORY-01", "剧情检索", "用户可按关键词、版本、世界、任务类型、系列和角色筛选剧情并查看场景。", "响应含总数、筛选项和详情；场景定位使用 mission_id/chunk_id。"),
        ("FR-ACT-01", "活动资料", "用户可浏览活动、详情和已发布玩家攻略。", "公开接口不返回未通过审核的内容。"),
        ("FR-CHAT-01", "智能问答", "用户提交消息和可选会话上下文，系统路由到专业 Agent 并返回统一 AI 响应。", "结果包含 answer、agent、协议、引用/校验/过滤信息和可选记忆建议。"),
        ("FR-CHAT-02", "会话持久化", "登录用户可查看会话、恢复消息并删除本人会话。", "资源按 user_id 隔离；删除级联清理消息。"),
        ("FR-CHAT-03", "异步聊天任务", "用户可创建、查询和重试长耗时聊天任务。", "状态至少覆盖 queued/running/succeeded/failed；重启后可恢复未完成任务。"),
        ("FR-MEM-01", "长期记忆", "系统只生成候选记忆，用户确认后方可创建；用户可查看、编辑、停用或删除。", "模型输出本身不得直接写库；所有记录属于当前用户。"),
        ("FR-PROFILE-01", "角色池", "用户可批量替换持有角色并标记喜爱角色。", "输入角色 ID 去重且必须存在；返回最新角色池。"),
        ("FR-PROG-01", "单角色养成", "用户可查看角色养成结构、计算等级与技能材料并保存练度。", "等级和技能区间通过校验；仅可保存已持有角色。"),
        ("FR-PROG-02", "多角色养成", "系统合并多个角色的材料需求并生成优先级建议。", "重复角色和非法区间返回 422；结果含汇总与证据版本。"),
        ("FR-PLAN-01", "养成方案", "用户可新增、查看、修改优先级/状态和删除养成方案。", "只操作本人方案；快照保留计算时输入与结果。"),
        ("FR-WEEK-01", "每周规划", "系统按启用方案、体力预算和周本剩余次数生成当前自然周任务。", "任务可勾选完成；生成结果记录证据版本和体力分配。"),
        ("FR-TEAM-01", "官方配队", "用户按核心角色、喜爱/排除角色、持有角色约束和生存位要求生成两套队伍。", "成员不重复并说明推荐依据；可区分理论队与可用队。"),
        ("FR-TEAM-02", "常用队伍", "登录用户可保存、查看、改名和删除四人队伍。", "队伍成员合法且位置唯一；按用户隔离。"),
        ("FR-CC-01", "自定义角色创作", "用户创建草稿并通过分阶段会话完善结构化字段。", "会话可恢复；确认阶段后推进；草稿与发布版本分离。"),
        ("FR-CC-02", "创作辅助", "系统可从对话抽取字段、生成头像提示或推荐自定义配队。", "无外部模型时仍允许固定提问和表单方式完成核心流程。"),
        ("FR-COM-01", "投稿与审核", "创作者提交版本后进入待审核；管理员批准或驳回。", "仅 approved/published 内容进入社区公开查询。"),
        ("FR-MOD-01", "社区治理", "管理员可查询已发布/已下架内容，并执行下架或恢复。", "操作记录内容类型、对象、原因、操作者和时间。"),
        ("FR-SYS-01", "健康与文档", "系统提供健康检查、Swagger UI 和 OpenAPI JSON。", "后端健康接口可被容器探针访问；OpenAPI 可导入 Apifox。"),
    ]
    add_table(
        doc,
        ["编号", "名称", "需求描述", "验收要点"],
        [list(row) for row in functional],
        [1250, 1450, 3600, 3060],
        font_size=8.8,
    )

    add_heading(doc, "5. 核心业务规则")
    add_bullets(
        doc,
        [
            "官方资料与玩家创作数据分离；玩家内容不得写入官方 Milvus collection。",
            "私有资源必须按当前 user_id 查询，避免通过猜测 UUID 访问他人数据。",
            "长期记忆必须经过用户确认；LLM 不能直接修改记忆、队伍、计划或公开状态。",
            "配队、养成、材料和每周规划中的数值结论优先由确定性工具计算并校验。",
            "AI 回答遵循 Claim → Citation → Validation → Filtering，无法获得可靠证据时应明确降级或拒答。",
            "投稿需经历 draft/submitted/approved 或 rejected 等状态；发布版本与当前草稿互不覆盖。",
        ],
    )

    add_heading(doc, "6. 非功能需求")
    nfrs = [
        ("NFR-PERF-01", "性能目标", "普通目录接口在目标环境 P95 ≤ 500 ms；纯数据库写操作 P95 ≤ 800 ms。", "使用固定数据集和并发模型压测；排除外部 LLM 接口。"),
        ("NFR-PERF-02", "AI 时延", "同步聊天应在 120 秒客户端超时内完成；长任务优先使用异步 job。", "记录路由、检索、重排、生成阶段耗时。"),
        ("NFR-CAP-01", "容量", "目标容量必须以用户数、日对话量、知识条目数和向量数定义。", "上线前形成容量基线和 70%/85% 告警阈值。"),
        ("NFR-AVL-01", "可用性", "单机试运行提供健康检查和自动重启；生产目标可用性由 SLA 另行约定。", "故障注入验证依赖不可用时的错误码和降级。"),
        ("NFR-REL-01", "一致性", "用户写操作事务化；外键资源删除按设计级联或受限。", "数据库约束、迁移测试和并发冲突测试。"),
        ("NFR-SEC-01", "认证安全", "密码仅保存强哈希，令牌有有效期，生产禁用默认 JWT 密钥。", "安全配置检查、401/403 越权测试、秘密扫描。"),
        ("NFR-SEC-02", "输入与权限", "所有请求经 Pydantic 校验；管理员操作需服务器端授权。", "非法枚举、超长文本、IDOR 和注入测试。"),
        ("NFR-PRIV-01", "隐私", "仅保存提供服务所需用户数据，支持删除会话、记忆和自有内容。", "隐私清单、保留周期和删除验证。"),
        ("NFR-MAIN-01", "可维护性", "前后端模块分层，数据库变更必须使用 Alembic。", "ruff、pytest、type-check 和 build 作为合并门禁。"),
        ("NFR-COMP-01", "兼容性", "支持现代 Chromium/Edge/Firefox；后端支持 Python 3.11–3.12，前端 Node.js 20+。", "浏览器矩阵与构建验证。"),
        ("NFR-OBS-01", "可观测性", "记录请求 ID、状态码、耗时、Agent 路由、依赖错误与模型成本，不记录密钥和明文密码。", "日志抽查、敏感信息检查、指标告警演练。"),
        ("NFR-DATA-01", "数据质量", "知识数据具备来源、版本、清洗与对齐报告；引用可回溯。", "运行现有数据审计脚本并检查差异。"),
        ("NFR-BACKUP-01", "备份恢复", "生产需定义 PostgreSQL 与 Milvus 的 RPO/RTO 并完成恢复演练。", "从备份恢复至隔离环境并验证关键查询。"),
        ("NFR-UX-01", "易用与无障碍", "关键流程支持键盘操作、明确错误提示和响应式布局，目标为 WCAG 2.1 AA。", "自动检查加人工键盘/读屏抽查。"),
    ]
    add_table(
        doc,
        ["编号", "类别", "需求/目标", "验证方式"],
        [list(row) for row in nfrs],
        [1350, 1250, 4100, 2660],
        font_size=8.8,
    )

    add_heading(doc, "7. 外部接口与依赖")
    add_table(
        doc,
        ["依赖", "用途", "失败影响", "降级/处置"],
        [
            ["DeepSeek / Qwen", "意图识别或有据生成", "AI 回答受限", "mock、规则路由、明确 503 或简化回答"],
            ["BGE 模型服务", "向量化与重排", "语义检索不可用或降质", "健康检查、CPU 模式、本地检索回退"],
            ["Milvus", "向量知识检索", "RAG 证据不可用", "结构化目录仍可用，禁止无依据编造"],
            ["PostgreSQL", "账号与业务状态", "登录和保存功能不可用", "快速失败、告警、从备份恢复"],
            ["本地 JSON/Markdown/图片", "官方目录与静态资源", "内容缺失或显示异常", "启动检查、清单校验、占位资源"],
        ],
        [1800, 2200, 2400, 2960],
    )

    add_heading(doc, "8. 验收建议")
    add_numbered(
        doc,
        [
            "按功能需求编号建立测试用例，至少覆盖成功、校验失败、未认证、越权和资源不存在。",
            "使用 OpenAPI/Apifox 建立接口回归集，登录后自动保存 token 并运行受保护接口。",
            "使用 pytest、ruff、前端 type-check 与 build 验证代码质量。",
            "在目标部署环境执行容量与时延测试，将测量结果回填非功能需求。",
            "完成数据库迁移、备份恢复、模型服务不可用和外部 LLM 超时演练。",
        ],
    )
    return save_doc(doc, "02_需求分析报告.docx")


def build_design() -> Path:
    doc = new_doc(
        "系统设计报告",
        "总体架构、数据边界与主要功能设计",
        "SRA-DES-001",
        "系统设计",
    )
    add_callout(
        doc,
        "设计主线",
        "采用前后端分离、领域服务分层、混合式 Multi-Agent 与双存储边界。"
        "生成式模型负责语义与表达，确定性服务负责数值计算和校验。",
    )
    add_heading(doc, "1. 设计目标与原则")
    add_bullets(
        doc,
        [
            "可追溯：AI 事实性回答携带引用、校验和过滤信息。",
            "可控：用户数据写入和内容发布必须由用户或管理员显式触发。",
            "职责分离：Vue 负责交互，FastAPI 编排业务，PostgreSQL 保存事务数据，Milvus 保存向量知识。",
            "可降级：外部模型或向量服务异常时，结构化目录和确定性能力尽可能保持可用。",
            "可部署：通过 Docker Compose 复现完整依赖，通过 Alembic 管理数据库演进。",
        ],
    )

    add_heading(doc, "2. 总体架构设计")
    add_table(
        doc,
        ["层次", "组件", "职责", "主要接口"],
        [
            ["表现层", "Vue 3 + TypeScript + Pinia + Router", "页面、表单、状态、鉴权令牌、错误展示", "HTTP/JSON"],
            ["接入层", "Nginx + FastAPI", "静态页面、REST 路由、校验、CORS、OpenAPI", "/api/v1/*"],
            ["应用层", "API Routes + Services", "用例编排、授权、事务边界、响应模型", "Python 调用"],
            ["智能体层", "黑塔主 Agent、Router、专业 Agents", "意图识别、上下文构造、工具选择、回答整合", "统一 AIResponse"],
            ["检索与模型层", "BGE 服务、Milvus、DeepSeek/Qwen", "嵌入、重排、证据检索与生成", "HTTP / Milvus SDK"],
            ["数据层", "PostgreSQL、JSON/Markdown、图片", "用户事务数据、官方知识、静态资源", "SQLAlchemy / 文件读取"],
        ],
        [1300, 2450, 3500, 2110],
    )
    add_heading(doc, "2.1 逻辑调用链", level=2)
    add_callout(
        doc,
        "请求流",
        "浏览器 → Nginx/Vue → FastAPI 路由 → 领域服务或黑塔主 Agent → "
        "PostgreSQL / Milvus / BGE / LLM → 统一 JSON 响应 → 前端呈现。",
        fill="F7F9FB",
    )
    add_heading(doc, "2.2 部署拓扑", level=2)
    add_table(
        doc,
        ["服务", "容器/端口", "依赖", "持久化"],
        [
            ["frontend", "宿主机 8080 → 80", "backend", "镜像内静态文件"],
            ["backend", "8000", "PostgreSQL、Milvus、BGE", "业务数据在 PostgreSQL"],
            ["bge-model-service", "8001", "本地模型目录、可选 GPU", "模型目录只读挂载"],
            ["postgres", "容器 5432", "无", "postgres_data"],
            ["milvus-standalone", "宿主机默认 19531 → 19530", "etcd、MinIO", "milvus_data"],
            ["etcd / minio", "内部端口", "无", "etcd_data / minio_data"],
        ],
        [1900, 2200, 2900, 2360],
    )

    add_heading(doc, "3. 模块设计")
    modules = [
        ("认证与用户域", "注册、登录、JWT、当前用户、管理员识别", "users", "auth 路由、security"),
        ("官方目录域", "角色/物品/光锥/遗器/怪物详情与关联跳转", "文件知识", "catalog service"),
        ("剧情与活动域", "剧情筛选、场景查看、活动与玩家攻略", "文件 + activity_guides", "story/activity services"),
        ("对话与记忆域", "同步/异步问答、会话恢复、用户确认记忆", "conversations、chat_messages、chat_jobs、user_memories", "chat/memory services"),
        ("角色养成域", "材料计算、练度保存、多角色汇总、每周计划", "user_character_progress、progression_plans、weekly_plans", "progression/planning services"),
        ("配队域", "官方队伍推荐、理论/可用队、常用队伍", "user_teams、user_team_members", "team services/agents"),
        ("创作与社区域", "分阶段创作、版本、投稿、审核、治理", "custom_*、moderation_actions", "custom character/review services"),
        ("RAG 与智能体域", "路由、检索、重排、证据协议、风格统一", "Milvus + 文件知识", "agents/rag/llm"),
    ]
    add_table(
        doc,
        ["模块", "主要职责", "数据", "实现边界"],
        [list(row) for row in modules],
        [1700, 2750, 3000, 1910],
        font_size=8.9,
    )

    add_heading(doc, "4. Multi-Agent 与 RAG 设计")
    add_heading(doc, "4.1 智能体协作", level=2)
    add_numbered(
        doc,
        [
            "编排层（黑塔主 Agent）：接收当前消息、最近会话、渐进式会话摘要、角色池和已确认记忆，"
            "生成规划轨迹事件、页面动作（深链与预填）与衍生问题建议，并按意图类型执行回答预算。",
            "注册表（AgentRegistry）：全部专业 Agent 单点注册，声明服务意图、能力描述、页面动作与"
            "兜底关键词；主 Agent 的自我介绍由注册表生成，新增 Agent 自动纳入能力自述。",
            "调度层（Router）：使用专用语义模型（glm-4.7 意图档，失败回退 4.6）识别意图、消解指代、"
            "将复合问题拆解为独立子任务；模型失败时按确定性关键词回退，仍失败交 RAG 兜底。",
            "执行层（专业 Agent）：调用结构化服务、确定性计算或 RAG 检索；创作自查、档案速查等"
            "原页面内能力已收编为一等 Agent 进入聊天链路。",
            "专业结果统一为主张、引用、校验、过滤与查询步骤；复合子任务由复核 Agent 合并并校验覆盖度。",
            "主 Agent 统一表达风格（风格池轮换避免模板化），但不得改变证据含义或新增无依据事实。",
        ],
    )
    add_heading(doc, "4.1.1 功能询问与自我认知", level=3)
    add_para(
        doc,
        "「你能帮我做什么」「XX 功能是干嘛的」类询问由能力注册表直接生成回答，"
        "不进入功能 Agent 误触发业务；注册表同时驱动页面动作与规划提示，保证自述与实现一致。"
    )
    add_heading(doc, "4.2 RAG 检索流水线", level=2)
    add_callout(
        doc,
        "RAG 流程",
        "独立问题 → BGE-M3 查询向量 → Milvus Top-K 候选 → 标题/关键词加权 → "
        "BGE Reranker → 生成 Citation ID → 有据生成 → 引用与置信度校验 → 过滤。",
        fill="F7F9FB",
    )
    add_para(
        doc,
        "该设计把“检索到什么”和“如何表达”分开。无引用或低置信主张在输出前被过滤；"
        "检索服务不可用时，系统应返回明确错误或改用结构化数据，而不是让模型凭空补全游戏事实。"
    )

    add_heading(doc, "5. 主要功能设计")
    add_heading(doc, "5.1 认证与授权", level=2)
    add_bullets(
        doc,
        [
            "注册时校验用户名和邮箱唯一性，密码经安全哈希后保存。",
            "登录成功签发有期限的 JWT；前端在请求拦截器加入 Authorization: Bearer <token>。",
            "依赖函数统一解析当前用户和管理员身份；私有查询同时限定资源 ID 与 user_id。",
        ],
    )
    add_heading(doc, "5.2 智能问答与异步任务", level=2)
    add_bullets(
        doc,
        [
            "同步接口适用于可在客户端 120 秒超时内完成的问答。",
            "异步 job 接口保存 request_payload、status、response_payload、error_message 与时间戳。",
            "应用启动时恢复 queued/running 任务；失败任务允许显式重试，避免前端无限等待。",
        ],
    )
    add_heading(doc, "5.3 养成与每周规划", level=2)
    add_bullets(
        doc,
        [
            "角色养成服务读取结构化材料规则，校验等级和技能区间后进行材料合并。",
            "ProgressionPlan 保存请求、材料和推荐快照，使后续数据版本变化时仍能解释历史结果。",
            "WeeklyPlan 以自然周、体力预算、周本次数和启用方案生成任务列表，并支持逐任务完成状态更新。",
        ],
    )
    add_heading(doc, "5.4 配队", level=2)
    add_bullets(
        doc,
        [
            "请求输入核心角色、偏好、排除、生存位、仅使用持有角色与游戏模式（混沌回忆/虚构叙事/末日幻影）等约束。",
            "候选生成：按槽位计划对全池角色六因子排序（定位/机制标签/官方共现/体系引擎亲和/增益价值适配/属性），"
            "12 次轮转保证多样性；同一切换单位（开拓者各形态、三月七双命途）的变体组互斥由候选过滤与最终校验双层拦截。",
            "官方存储队整队注入为显式候选，并按体系引擎件生成等价类变体（替换位优先满足核心未满足硬需求）；"
            "核心首选存储队保证出现在理论推荐中，缺失角色照常标注。",
            "机制模拟重打分（scoring v3）：知识分 55% + 场景 20% + 机械核心 15% + 数据置信 10%；"
            "知识分基于机制画像逐条判定核心硬/软需求满足度、引擎参与度与体系增益价值适配。",
            "机制画像与受控词表存储于 docs/team_knowledge（LLM 从技能文本标注 + 人工修订），"
            "校准测试保证官方存储队得分高于体系外扰动版，评分规则改动需先通过考卷。",
            "LLM 终审只输出解说与结论标签，不能修改成员、分数或顺序；解说必须引用成员机制引擎与需求判定。",
            "保存队伍时使用成员表记录 position，并通过唯一约束防止重复位置或角色。",
        ],
    )
    add_heading(doc, "5.5 自定义角色与审核", level=2)
    add_bullets(
        doc,
        [
            "CustomCharacter 保存作品身份与可见性；CustomCharacterVersion 保存不可混淆的版本 payload。",
            "CreatorSession 与 Message 保存四阶段引导上下文和结构化抽取结果，支持恢复。",
            "提交、批准、驳回、下架和恢复均由显式接口驱动；发布版本与草稿指针分离。",
            "ModerationAction 记录内容类型、对象、动作、原因、操作者和详情，支持治理追溯。",
        ],
    )

    add_heading(doc, "6. 数据设计边界")
    add_table(
        doc,
        ["数据类别", "存储", "访问方式", "边界规则"],
        [
            ["用户事务数据", "PostgreSQL", "SQLAlchemy", "user_id 隔离、事务与外键约束"],
            ["官方结构化知识", "JSON/Markdown", "Catalog/Story/Activity 服务", "只读挂载、版本化清洗"],
            ["语义向量", "Milvus", "RAG Retriever", "仅官方知识 collection"],
            ["图片资源", "docs/data/assets", "FastAPI StaticFiles/Nginx", "清单校验与只读挂载"],
            ["本地模型权重", "models 目录", "BGE 服务", "容器只读，不写入镜像"],
        ],
        [1750, 1850, 2300, 3460],
    )

    add_heading(doc, "7. 接口与错误设计")
    add_bullets(
        doc,
        [
            "REST API 统一使用 /api/v1 前缀，JSON 为主要媒体类型。",
            "Pydantic 负责请求/响应契约；FastAPI 生成 OpenAPI，并通过 /docs 提供 Swagger UI。",
            "典型错误：401 未认证、403 权限不足、404 不存在或不属于当前用户、409 唯一性冲突、422 业务/字段校验、503 依赖不可用。",
            "对外错误信息应可行动但不泄露堆栈、密码哈希、API Key、数据库连接串或内部文件路径。",
        ],
    )

    add_heading(doc, "8. 安全设计")
    add_table(
        doc,
        ["威胁", "控制", "验证"],
        [
            ["账号接管", "密码哈希、JWT 有效期、生产强密钥", "认证失败、过期令牌和弱配置测试"],
            ["越权访问", "服务器端 current_user/admin 依赖、user_id 限定", "IDOR 与角色权限测试"],
            ["提示注入/幻觉", "工具白名单、证据协议、过滤与数据边界", "恶意提示、无证据问题和依赖故障测试"],
            ["敏感信息泄露", ".env 忽略、日志脱敏、错误响应收敛", "秘密扫描与日志抽查"],
            ["社区违规内容", "提交审核、管理员治理与操作记录", "状态机和审计完整性测试"],
        ],
        [1900, 4300, 3160],
    )

    add_heading(doc, "9. 运行与扩展设计")
    add_bullets(
        doc,
        [
            "当前 Compose 适合开发和单机演示；生产可将 PostgreSQL、对象存储和向量库替换为托管或高可用部署。",
            "模型服务已独立，可按 GPU 容量横向扩展；后端需避免在进程内保存不可恢复的关键状态。",
            "建议为请求、Agent 步骤、外部模型、BGE 与 Milvus 建立统一 trace_id 和指标。",
            "数据库结构通过 Alembic 向前迁移；发布流程应先备份、再迁移、再滚动更新并保留回滚方案。",
        ],
    )
    return save_doc(doc, "03_系统设计报告.docx")


def schema_label(schema: Any) -> str:
    if not schema:
        return "无"
    if "$ref" in schema:
        return schema["$ref"].split("/")[-1]
    if "allOf" in schema:
        return " + ".join(schema_label(x) for x in schema["allOf"])
    if schema.get("type") == "array":
        return f"array<{schema_label(schema.get('items', {}))}>"
    return schema.get("title") or schema.get("type") or "内联结构"


def operation_auth(operation: dict[str, Any]) -> str:
    security = operation.get("security")
    if security == []:
        return "公开"
    if security:
        return "Bearer JWT"
    return "按路由依赖（多数私有接口需 Bearer JWT）"


def build_api() -> tuple[Path, Path]:
    spec = app.openapi()
    spec.setdefault("servers", [{"url": "http://localhost:8000", "description": "本地 FastAPI"}])
    openapi_path = OUTPUT_DIR / "star-rail-agents.openapi.json"
    openapi_path.write_text(json.dumps(spec, ensure_ascii=False, indent=2), encoding="utf-8")

    doc = new_doc(
        "接口文档",
        "REST API 功能说明与 Apifox 导入/查看指南",
        "SRA-API-001",
        "接口文档",
    )
    add_callout(
        doc,
        "接口基线",
        f"本文档从 FastAPI OpenAPI 自动生成，共 {len(spec.get('paths', {}))} 条路径、"
        f"{sum(len([m for m in item if m in {'get','post','put','patch','delete'}]) for item in spec.get('paths', {}).values())} 个操作。"
        "实际调试以运行中服务的 /openapi.json 为准。",
    )
    add_heading(doc, "1. 基本约定")
    add_table(
        doc,
        ["项目", "说明"],
        [
            ["本地 API 根地址", "http://localhost:8000"],
            ["业务前缀", "/api/v1"],
            ["交互式文档", "http://localhost:8000/docs"],
            ["OpenAPI JSON", "http://localhost:8000/openapi.json"],
            ["认证", "Authorization: Bearer <access_token>"],
            ["请求/响应格式", "application/json；文件/静态资源接口除外"],
            ["超时", "前端普通请求 30 秒；聊天/配队等长任务可到 120 秒或使用异步 job"],
        ],
        [2500, 6860],
    )
    add_heading(doc, "2. 在 Apifox 中导入与查看")
    add_numbered(
        doc,
        [
            "启动服务：在项目根目录执行 docker compose up -d --build，并确认 backend 健康。",
            "打开浏览器检查 http://localhost:8000/openapi.json 能返回 JSON；也可直接使用随本文档交付的 star-rail-agents.openapi.json。",
            "在 Apifox 新建项目，选择“导入数据”或“项目设置 → 导入数据”，数据格式选择 OpenAPI/Swagger。",
            "在线导入填写 http://localhost:8000/openapi.json；离线导入选择 star-rail-agents.openapi.json。",
            "创建环境变量 baseUrl=http://localhost:8000，并将接口前缀保留在路径中。",
            "先调用 POST /api/v1/auth/register 或 /auth/login，复制响应 access_token。",
            "在项目或目录的 Auth 中选择 Bearer Token，值设置为 token；也可在登录接口后置脚本中把 token 写入环境变量。",
            "进入任一接口可查看说明、参数、Body Schema、响应模型和状态码；切换“运行”页填写示例并发送请求。",
        ],
    )
    add_heading(doc, "2.1 Apifox 登录后置脚本示例", level=2)
    add_callout(
        doc,
        "脚本",
        "const body = pm.response.json(); if (body.access_token) { "
        "pm.environment.set('access_token', body.access_token); }",
        fill="F7F9FB",
    )
    add_para(
        doc,
        "随后将项目级 Bearer Token 设置为 {{access_token}}。若 Apifox 版本使用不同脚本对象，"
        "可直接复制 token 到环境变量；核心要求是最终请求头为 Authorization: Bearer <token>。"
    )
    add_heading(doc, "3. 常见状态码")
    add_table(
        doc,
        ["状态码", "含义", "处理建议"],
        [
            ["200/201", "读取/创建成功", "按响应模型继续业务流程"],
            ["204", "删除成功且无响应体", "客户端清理本地状态"],
            ["401", "未登录、令牌缺失或失效", "重新登录并更新 Bearer token"],
            ["403", "当前账号无管理员权限", "切换授权账号或停止调用"],
            ["404", "资源不存在或不属于当前用户", "检查 ID 与资源所有权"],
            ["409", "用户名、邮箱或业务唯一项冲突", "修改输入或查询现有记录"],
            ["422", "字段或业务规则校验失败", "查看 detail 并修正类型、范围或枚举"],
            ["503", "模型、向量库等依赖不可用", "检查健康状态，稍后重试或使用降级能力"],
        ],
        [1200, 3600, 4560],
    )

    grouped: dict[str, list[tuple[str, str, dict[str, Any]]]] = defaultdict(list)
    allowed_methods = {"get", "post", "put", "patch", "delete"}
    for path, item in spec.get("paths", {}).items():
        for method, operation in item.items():
            if method not in allowed_methods:
                continue
            tags = operation.get("tags") or ["未分组"]
            grouped[tags[0]].append((method.upper(), path, operation))

    add_heading(doc, "4. 接口总览")
    overview_rows = []
    for tag, operations in sorted(grouped.items()):
        for method, path, operation in sorted(operations, key=lambda x: (x[1], x[0])):
            overview_rows.append(
                [
                    tag,
                    method,
                    path,
                    operation.get("summary") or operation.get("operationId") or "接口操作",
                    operation_auth(operation),
                ]
            )
    add_table(
        doc,
        ["分组", "方法", "路径", "功能", "认证"],
        overview_rows,
        [1200, 850, 3450, 2760, 1100],
        font_size=8.1,
    )

    add_heading(doc, "5. 接口明细")
    section_number = 1
    for tag, operations in sorted(grouped.items()):
        add_heading(doc, f"5.{section_number} {tag}", level=2)
        section_number += 1
        for method, path, operation in sorted(operations, key=lambda x: (x[1], x[0])):
            summary = operation.get("summary") or operation.get("operationId") or "接口操作"
            add_heading(doc, f"{method} {path}", level=3)
            add_para(doc, f"功能：{summary}")
            if operation.get("description"):
                add_para(doc, f"说明：{operation['description']}")
            request_schema = "无"
            content = operation.get("requestBody", {}).get("content", {})
            if content:
                media = content.get("application/json") or next(iter(content.values()))
                request_schema = schema_label(media.get("schema", {}))
            parameters = operation.get("parameters", [])
            parameter_text = "无"
            if parameters:
                parts = []
                for param in parameters:
                    required = "必填" if param.get("required") else "可选"
                    parts.append(
                        f"{param.get('name')}({param.get('in')}, {required}, "
                        f"{schema_label(param.get('schema', {}))})"
                    )
                parameter_text = "；".join(parts)
            responses = operation.get("responses", {})
            response_text = "；".join(
                f"{code}: {schema_label((resp.get('content', {}).get('application/json') or {}).get('schema', {}))}"
                for code, resp in responses.items()
            )
            add_table(
                doc,
                ["项目", "内容"],
                [
                    ["认证", operation_auth(operation)],
                    ["参数", parameter_text],
                    ["请求体模型", request_schema],
                    ["响应", response_text or "见 OpenAPI"],
                    ["Operation ID", operation.get("operationId", "—")],
                ],
                [1900, 7460],
                font_size=8.7,
            )

    add_heading(doc, "6. 模型服务内部接口")
    add_para(
        doc,
        "BGE 模型服务默认监听 8001，主要供后端容器调用，不建议直接暴露给公网。其核心接口如下："
    )
    add_table(
        doc,
        ["方法与路径", "功能", "请求", "响应/说明"],
        [
            ["GET /health", "检查模型文件、设备和加载状态", "无", "健康信息；供容器探针使用"],
            ["POST /embeddings", "批量生成 BGE-M3 向量", "texts: 1–128 条文本", "model、dimension、vectors"],
            ["POST /rerank", "对候选段落重排", "query + 1–100 passages", "model、scores"],
        ],
        [2200, 2500, 2300, 2360],
    )
    add_heading(doc, "7. 接口调试注意事项")
    add_bullets(
        doc,
        [
            "OpenAPI 文件是静态快照；代码发生路由或 Schema 变更后，应重新导入运行中服务的 /openapi.json。",
            "静态资源和个别公开接口不需要令牌；私有与管理员接口仍以服务器端依赖校验为准。",
            "调用聊天或配队接口时提高客户端超时；需要可靠轮询时使用 /chat/jobs。",
            "不要在截图、日志、测试数据或共享 Apifox 环境中保存真实 API Key、密码或长期有效 token。",
        ],
    )
    return save_doc(doc, "04_接口文档_Apifox使用指南.docx"), openapi_path


def type_label(column) -> str:
    text = str(column.type)
    if not column.nullable:
        text += " NOT NULL"
    return text


def column_default(column) -> str:
    if column.default is not None:
        return str(column.default.arg)
    if column.server_default is not None:
        return str(column.server_default.arg)
    return "—"


def table_relations(table) -> str:
    rels = []
    for fk in table.foreign_keys:
        rels.append(f"{fk.parent.name} → {fk.column.table.name}.{fk.column.name}")
    return "；".join(rels) if rels else "无外键"


def build_database() -> Path:
    doc = new_doc(
        "数据库设计文档",
        "数据模型、使用说明、迁移与部署步骤",
        "SRA-DB-001",
        "数据库设计",
    )
    metadata = Base.metadata
    tables = sorted(metadata.tables.values(), key=lambda t: t.name)
    add_callout(
        doc,
        "数据基线",
        f"业务数据库使用 PostgreSQL 16，当前 ORM 元数据包含 {len(tables)} 张表；"
        "Milvus 用于官方知识向量检索，不替代 PostgreSQL 的事务数据。",
    )
    add_heading(doc, "1. 数据库总体设计")
    add_table(
        doc,
        ["数据系统", "保存内容", "一致性特征", "持久化"],
        [
            ["PostgreSQL 16", "用户、角色池、记忆、会话、任务、队伍、计划、创作与审核", "事务、外键、唯一约束、索引", "postgres_data 命名卷"],
            ["Milvus 2.5.2", "官方知识文本向量及检索元数据", "collection 级管理、最终重建可恢复", "milvus_data + etcd + MinIO"],
            ["JSON/Markdown", "官方角色、物品、剧情、活动等源数据", "只读、版本化、可审计", "docs 目录只读挂载"],
            ["图片/模型文件", "静态图片、BGE 权重", "清单校验/只读", "宿主机目录挂载"],
        ],
        [1800, 3350, 2300, 1910],
    )
    add_heading(doc, "2. PostgreSQL 逻辑关系")
    add_bullets(
        doc,
        [
            "users 为用户域根实体，关联角色池、记忆、练度、队伍、会话、聊天任务、养成方案、每周计划和创作内容。",
            "user_teams 与 user_team_members 为一对多；团队删除时成员级联删除。",
            "conversations 与 chat_messages 为一对多；会话删除时消息级联删除。",
            "custom_characters 拥有多个版本和创作会话；会话拥有多条消息。",
            "activity_guides 与 moderation_actions 保存社区内容状态和治理记录。",
            "官方角色等目录数据不复制到业务表，业务表通过 character_id 等稳定标识引用文件知识。",
        ],
    )

    add_heading(doc, "3. 表清单")
    summary_rows = []
    for table in tables:
        pk = ", ".join(col.name for col in table.primary_key.columns)
        summary_rows.append(
            [table.name, len(table.columns), pk or "—", table_relations(table)]
        )
    add_table(
        doc,
        ["表名", "字段数", "主键", "外键关系"],
        summary_rows,
        [2600, 950, 1700, 4110],
        font_size=8.7,
    )

    add_heading(doc, "4. 数据字典")
    for idx, table in enumerate(tables, start=1):
        add_heading(doc, f"4.{idx} {table.name}", level=2)
        rows = []
        for column in table.columns:
            flags = []
            if column.primary_key:
                flags.append("PK")
            if column.unique:
                flags.append("UNIQUE")
            if column.index:
                flags.append("INDEX")
            if column.foreign_keys:
                flags.append("FK")
            rows.append(
                [
                    column.name,
                    type_label(column),
                    column_default(column),
                    "、".join(flags) if flags else "—",
                    "；".join(str(fk.target_fullname) for fk in column.foreign_keys) or "—",
                ]
            )
        add_table(
            doc,
            ["字段", "类型/空值", "默认值", "约束/索引", "引用"],
            rows,
            [1900, 2600, 1700, 1450, 1710],
            font_size=8.2,
        )
        constraints = []
        for constraint in table.constraints:
            name = constraint.name or constraint.__class__.__name__
            cols = ", ".join(col.name for col in getattr(constraint, "columns", []))
            if constraint.__class__.__name__ not in {"PrimaryKeyConstraint", "ForeignKeyConstraint"}:
                constraints.append(f"{name}({cols})")
        if constraints:
            add_para(doc, "其他约束：" + "；".join(constraints))

    add_heading(doc, "5. 索引与一致性设计")
    add_bullets(
        doc,
        [
            "用户名称与邮箱使用唯一索引，防止重复账号。",
            "user_id、character_id、status、week_start 等高频筛选字段建立索引。",
            "队伍成员通过 team_id + character_id、team_id + position 等组合唯一性保证成员和位置不重复。",
            "JSON/JSONB 字段保存请求和结果快照，适合结构演进；如需高频按内部属性检索，应再评估 GIN 索引和字段规范化。",
            "跨存储一致性不采用分布式事务；官方向量库可从源文档重新构建，PostgreSQL 业务数据必须备份。",
        ],
    )

    add_heading(doc, "6. 使用说明")
    add_heading(doc, "6.1 连接配置", level=2)
    add_callout(
        doc,
        "Docker 连接串",
        "postgresql+psycopg://star_rail:<password>@postgres:5432/star_rail_agents",
        fill="F7F9FB",
    )
    add_para(
        doc,
        "后端通过 DATABASE_URL 读取连接。宿主机直接调试后端时，主机名应改为 localhost，"
        "并确保 PostgreSQL 端口已映射或使用 IDE 的 Docker 服务网络。生产环境不得使用 change-me。"
    )
    add_heading(doc, "6.2 ORM 与会话", level=2)
    add_bullets(
        doc,
        [
            "SQLAlchemy DeclarativeBase 定义模型，SessionLocal 创建同步会话。",
            "FastAPI 依赖 get_db 在请求结束时关闭会话；写操作应在服务层明确 commit/rollback。",
            "数据库模型变更不得使用 create_all 代替迁移，应新增 Alembic revision 并评审升级/回滚逻辑。",
        ],
    )
    add_heading(doc, "6.3 向量库使用", level=2)
    add_bullets(
        doc,
        [
            "默认 collection 为 star_rail_knowledge_bge_m3_v2，可通过 MILVUS_COLLECTION 覆盖。",
            "BGE-M3 生成查询和文档向量，reranker 对候选段落重排。",
            "玩家自定义角色和社区内容不得写入官方 collection；如未来需要检索，必须建立独立 collection 与权限策略。",
        ],
    )

    add_heading(doc, "7. 部署步骤")
    add_numbered(
        doc,
        [
            "安装 Docker Desktop；GPU 模式另行安装 NVIDIA 驱动与 Container Toolkit。",
            "复制 .env.example 为 .env，修改 POSTGRES_PASSWORD、JWT_SECRET_KEY、管理员账号和模型配置。",
            "准备 models/bge-m3、models/bge-reranker-v2-m3 与 docs/data/assets；确认模型 config 和权重完整。",
            "在项目根目录执行 docker compose up -d --build。",
            "执行 docker compose ps，等待 postgres、milvus-standalone、bge-model-service 和 backend 健康。",
            "打开 http://localhost:8000/api/v1/health 验证后端；打开 /docs 验证 OpenAPI。",
            "后端容器启动命令会先执行 alembic upgrade head；在空库上确认所有迁移成功。",
        ],
    )
    add_heading(doc, "7.1 本地开发迁移", level=2)
    add_callout(
        doc,
        "PowerShell",
        "cd backend；.venv\\Scripts\\python -m alembic upgrade head；"
        ".venv\\Scripts\\python -m uvicorn app.main:app --reload",
        fill="F7F9FB",
    )
    add_heading(doc, "7.2 新增迁移", level=2)
    add_numbered(
        doc,
        [
            "修改 SQLAlchemy 模型并评审约束、默认值、数据兼容和索引。",
            "在 backend 目录生成迁移：alembic revision --autogenerate -m \"change description\"。",
            "人工检查 upgrade 和 downgrade，特别关注非空字段、枚举/状态和大表索引。",
            "在备份或临时数据库执行 upgrade head，再运行测试。",
            "生产发布前备份；先迁移数据库，再发布兼容新旧结构的应用版本。",
        ],
    )

    add_heading(doc, "8. 备份与恢复")
    add_table(
        doc,
        ["对象", "备份建议", "恢复验证"],
        [
            ["PostgreSQL", "定期 pg_dump 自定义格式；重要发布前额外备份", "恢复到隔离库，核对表数、迁移版本与关键用户流程"],
            ["Milvus", "按官方备份工具/存储快照策略备份；保留源文档与建库脚本", "恢复 collection 或从源文档重建并运行检索基准"],
            ["MinIO/etcd", "与 Milvus 版本一致地快照或备份", "在隔离环境验证 Milvus 元数据和对象完整"],
            ["配置", "保存脱敏配置模板与密钥清单，不备份明文密钥到代码库", "演练重新注入密钥和环境配置"],
        ],
        [1700, 4200, 3460],
    )
    add_callout(
        doc,
        "重要警告",
        "docker compose down -v 会删除命名卷并清空 PostgreSQL、Milvus、MinIO 和 etcd 数据。"
        "除非明确执行全量重置且已验证备份，否则不要使用 -v。",
        fill="FFF2CC",
    )

    add_heading(doc, "9. 安全与运维要求")
    add_bullets(
        doc,
        [
            "生产使用专用数据库账号和最小权限；限制 PostgreSQL、Milvus、MinIO、etcd 仅在受信网络访问。",
            "数据库密码、JWT 密钥和外部模型 API Key 进入密钥管理，不写入镜像、Git、日志或 Apifox 共享环境。",
            "启用连接数、慢查询、磁盘、缓存命中、锁等待、备份成功率和恢复时间监控。",
            "迁移需保留审计记录；破坏性迁移采用扩展—迁移数据—收缩的分阶段策略。",
            "明确数据保留与删除策略，尤其是会话、记忆、社区内容和审核记录。",
        ],
    )
    return save_doc(doc, "05_数据库设计与部署文档.docx")


def main() -> None:
    paths: list[Path] = []
    paths.append(build_feasibility())
    paths.append(build_requirements())
    paths.append(build_design())
    api_doc, openapi_path = build_api()
    paths.extend([api_doc, openapi_path])
    paths.append(build_database())
    print("\n".join(str(path) for path in paths))


if __name__ == "__main__":
    main()
